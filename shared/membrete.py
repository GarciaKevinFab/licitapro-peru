"""La hoja membretada de cada empresa, aplicada a los documentos que firma.

QUE ES

  El Word que cada empresa usa como papel: logo, razon social y RUC arriba,
  direccion y telefonos abajo, a veces una marca de agua. Las cotizaciones,
  cartas y declaraciones propias se presentan impresas en esa hoja; un PDF
  con el logo dibujado por nosotros no se parece a lo que la empresa manda
  cuando la hace a mano, y la entidad lo nota.

QUE DOCUMENTOS LA LLEVAN Y CUALES NO

  La llevan los documentos que escribe la empresa: la cotizacion por items,
  la declaracion jurada propia, la carta de presentacion, el pacto de
  integridad, la propuesta tecnica y economica. NO la llevan los formatos de
  la entidad que se llenan sobre su propio PDF (la SC del GOREMAD, la del
  SIGA, los anexos de la MPT): ese papel es de la entidad.

COMO SE APLICA

  Se guardan tres archivos por empresa en MEMBRETES_DIR, por convencion de
  nombre como la ficha RUC (shared/ficha_ruc.py), sin columna en la base:

    empresa_<id>.docx   el original, base de los documentos en Word
    empresa_<id>.pdf    su primera hoja convertida con LibreOffice, el fondo
                        de los documentos en PDF
    empresa_<id>.json   cuanto ocupan arriba la cabecera y abajo el pie

  La conversion se hace UNA vez, al subirlo: LibreOffice tarda segundos y no
  se le puede pedir en cada documento. Las bandas se miden sobre la pagina
  renderizada y no sobre los margenes del Word, porque los Word membretados
  mienten: el de C&M San Jose declara margen izquierdo 0 y el de Soluciones
  Informaticas margen inferior 0 con un pie de tres lineas.
"""
from __future__ import annotations

import asyncio
import io
import json
import logging
import os
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

log = logging.getLogger("shared.membrete")

# Anclado a la raiz del proyecto: "data/membretes" relativo dependia del
# directorio desde el que se arrancara el proceso. MEMBRETES_DIR sigue mandando.
DIR = Path(os.getenv("MEMBRETES_DIR")
           or Path(__file__).resolve().parents[1] / "data" / "membretes")
TAMANO_MAXIMO = 10 * 1024 * 1024
# Lo que puede ocupar el Word DESCOMPRIMIDO. Un .docx es un zip: 10 MB en el
# disco pueden ser gigas al abrirlo (la bomba zip), y quien lo abre es
# LibreOffice, en el servidor. Una hoja membretada con su logo no llega a 5 MB.
DESCOMPRIMIDO_MAXIMO = 50 * 1024 * 1024
MAX_PARTES = 2000

# Una fila de pixeles cuenta como "con tinta" si algun pixel baja de este
# gris. 235 deja fuera el antialias del papel blanco pero no un texto claro.
_UMBRAL = 235
# Un hueco de esta altura sin tinta termina la cabecera (o el pie). Los
# renglones de una misma cabecera quedan a menos; la marca de agua del centro
# de la hoja, mucho mas lejos (mas de 100 pt en la de Soluciones Informaticas).
#
# Era 18 pt y se quedaba corto: una cabecera con el logo y los datos en una
# tabla deja unos 30 pt hasta el filete que la subraya, la medicion cortaba
# antes del filete y el titulo del documento salia encima de la linea
# (2026-10-17, con las hojas rehechas de K&A y Minimarket San Juan).
_HUECO_PT = 40
# La cabecera y el pie tienen que EMPEZAR en esta fraccion de la hoja, desde
# su borde. Lo que empieza mas adentro es marca de agua o cuerpo.
_ALCANCE = 0.35
# Aire entre el membrete y el texto.
AIRE_PT = 14


class MembreteInvalido(Exception):
    """El mensaje esta escrito para el usuario."""


def _rutas(empresa_id: int) -> tuple[Path, Path, Path]:
    base = DIR / f"empresa_{int(empresa_id)}"
    return base.with_suffix(".docx"), base.with_suffix(".pdf"), base.with_suffix(".json")


def _escribir_atomico(destino: Path, contenido: bytes) -> None:
    """Escribe en un temporal del mismo directorio y lo pone en su sitio de golpe.

    POR QUE: escribir directo sobre el archivo deja, si el proceso muere a
    mitad, un PDF o un JSON cortado que despues se usa como bueno. Con
    `os.replace` el archivo es el viejo o el nuevo, nunca uno a medias.
    """
    destino.parent.mkdir(parents=True, exist_ok=True)
    fd, temporal = tempfile.mkstemp(dir=destino.parent, prefix=".tmp_",
                                    suffix=destino.suffix)
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(contenido)
        os.replace(temporal, destino)
    except BaseException:
        Path(temporal).unlink(missing_ok=True)
        raise


def de(empresa_id: int) -> dict | None:
    """{docx, pdf, cabecera_pt, pie_pt} si la empresa tiene membrete; si no, None."""
    docx, pdf, medidas = _rutas(empresa_id)
    if not (docx.is_file() and pdf.is_file()):
        return None
    try:
        bandas = json.loads(medidas.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        # Copiado a mano sin el .json: se mide ahora y se deja guardado.
        bandas = medir(pdf)
        _escribir_atomico(medidas, json.dumps(bandas).encode("utf-8"))
    return {"docx": docx, "pdf": pdf, **bandas}


# ─── Medir la cabecera y el pie ─────────────────────────────

def banda(filas: list[bool], alcance: int, hueco: int) -> int:
    """Filas que ocupa el membrete contando desde el borde que recibe `filas`.

    `filas` va del borde hacia el centro: True = la fila tiene tinta. La banda
    tiene que empezar antes de `alcance` y termina en el primer tramo de
    `hueco` filas en blanco. 0 si no hay membrete en ese borde.
    """
    ultima = None
    blancas = 0
    for i, tinta in enumerate(filas):
        if ultima is None:
            if i >= alcance:
                return 0
            if tinta:
                ultima = i
            continue
        if tinta:
            ultima = i
            blancas = 0
        else:
            blancas += 1
            if blancas >= hueco:
                break
    return 0 if ultima is None else ultima + 1


def medir(ruta_pdf: Path | str) -> dict:
    """Alto de la cabecera y del pie, en puntos, de la primera hoja."""
    import pypdfium2 as pdfium

    documento = pdfium.PdfDocument(str(ruta_pdf))
    try:
        # Escala 1 = 72 ppp: un pixel por punto basta para medir.
        imagen = documento[0].render(scale=1, grayscale=True).to_pil().convert("L")
    finally:
        documento.close()
    ancho, alto = imagen.size
    datos = imagen.tobytes()
    tinta = [min(datos[f * ancho:(f + 1) * ancho]) < _UMBRAL for f in range(alto)]
    alcance = int(alto * _ALCANCE)
    return {"cabecera_pt": float(banda(tinta, alcance, _HUECO_PT)),
            "pie_pt": float(banda(tinta[::-1], alcance, _HUECO_PT))}


# ─── Guardar y quitar ───────────────────────────────────────

# Cada <Relationship .../> de las partes .rels, y como reconocer las externas.
_RE_RELACION = re.compile(rb"<Relationship\b[^>]*>", re.IGNORECASE)
_RE_MODO_EXTERNO = re.compile(rb"TargetMode\s*=\s*[\"']External[\"']", re.IGNORECASE)
# Las externas que SI se aceptan: un hipervinculo es solo un enlace en el texto
# (la web de la empresa en el pie); nadie lo sigue al convertir.
_RE_HIPERVINCULO = re.compile(rb"Type\s*=\s*[\"'][^\"']*/hyperlink[\"']", re.IGNORECASE)

_NO_ES_WORD = ("La hoja membretada tiene que ser un Word (.docx). Si la tienes en "
               ".doc, ábrela en Word y guárdala como .docx.")


def _validar_docx(contenido: bytes) -> None:
    """Que sea un Word, que no sea una bomba zip y que no tire de la red.

    RELACIONES EXTERNAS: un .docx puede declarar que una imagen, una plantilla
    o un objeto incrustado viven en una URL (TargetMode="External"). Al
    convertirlo, LibreOffice -- en NUESTRO servidor -- iria a buscarlos: el
    usuario elegiria a que direccion se conecta la maquina (SSRF), incluidas
    las de la red interna. Se rechaza cualquier relacion externa salvo los
    hipervinculos, que solo son texto.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            partes = z.infolist()
            if "word/document.xml" not in {p.filename for p in partes}:
                raise MembreteInvalido(_NO_ES_WORD)
            if len(partes) > MAX_PARTES:
                raise MembreteInvalido("Ese Word tiene demasiadas partes internas.")
            if sum(p.file_size for p in partes) > DESCOMPRIMIDO_MAXIMO:
                raise MembreteInvalido(
                    "Ese Word ocupa demasiado al abrirlo (más de 50 MB). "
                    "Reduce las imágenes y vuelve a subirlo.")
            for parte in partes:
                if not parte.filename.lower().endswith(".rels"):
                    continue
                # El tamano declarado ya se comprobo arriba, y zipfile no
                # entrega mas bytes que los declarados.
                for rel in _RE_RELACION.findall(z.read(parte)):
                    if _RE_MODO_EXTERNO.search(rel) and not _RE_HIPERVINCULO.search(rel):
                        log.warning("Membrete rechazado, relacion externa: %r", rel[:200])
                        raise MembreteInvalido(
                            "Ese Word enlaza imágenes u objetos que están en "
                            "internet. Inserta las imágenes dentro del documento "
                            "(no vinculadas) y vuelve a subirlo.")
    except zipfile.BadZipFile as e:
        raise MembreteInvalido(_NO_ES_WORD) from e
    except (OSError, ValueError, RuntimeError, NotImplementedError) as e:
        # Zip cifrado, compresion desconocida o una parte danada.
        raise MembreteInvalido("No se pudo abrir ese Word; puede estar dañado.") from e


def _es_docx(contenido: bytes) -> bool:
    try:
        _validar_docx(contenido)
    except MembreteInvalido:
        return False
    return True


async def guardar(empresa_id: int, contenido: bytes) -> dict:
    """Guarda el Word, lo convierte a PDF y mide sus bandas.

    Todo se prepara en una carpeta temporal propia de esta subida y solo al
    final se pone en su sitio. POR QUE: el temporal tenia nombre fijo por
    empresa, asi que dos subidas a la vez se pisaban; y un fallo a mitad podia
    dejar el Word nuevo con el PDF viejo. El trabajo pesado (leer el zip,
    recortar el PDF, renderizarlo para medir) va en un hilo: en el bucle de
    eventos congelaba la web entera mientras duraba.
    """
    from shared.pdf_firma import docx_a_pdf

    if not contenido:
        raise MembreteInvalido("El archivo está vacío.")
    if len(contenido) > TAMANO_MAXIMO:
        raise MembreteInvalido("La hoja membretada pesa más de 10 MB.")
    await asyncio.to_thread(_validar_docx, contenido)

    docx, pdf, medidas = _rutas(empresa_id)
    DIR.mkdir(parents=True, exist_ok=True)
    carpeta = Path(await asyncio.to_thread(
        tempfile.mkdtemp, prefix=f".subiendo_{int(empresa_id)}_", dir=DIR))
    try:
        temporal = carpeta / docx.name
        await asyncio.to_thread(temporal.write_bytes, contenido)
        convertido, mensaje = await docx_a_pdf(str(temporal), str(carpeta))
        if not convertido:
            raise MembreteInvalido(f"No se pudo convertir la hoja a PDF. {mensaje}")
        pdf_nuevo = carpeta / f"fondo_{pdf.name}"
        bandas = await asyncio.to_thread(_preparar_fondo, Path(convertido), pdf_nuevo)
        await asyncio.to_thread(_colocar, temporal, pdf_nuevo, bandas, docx, pdf, medidas)
    finally:
        await asyncio.to_thread(shutil.rmtree, carpeta, True)

    log.info("Membrete de la empresa %s: %s", empresa_id, bandas)
    return bandas


def _preparar_fondo(convertido: Path, destino: Path) -> dict:
    _primera_hoja(convertido, destino)
    return medir(destino)


def _colocar(docx_nuevo: Path, pdf_nuevo: Path, bandas: dict,
             docx: Path, pdf: Path, medidas: Path) -> None:
    """Pone los tres archivos en su sitio, cada uno con un reemplazo atomico."""
    os.replace(pdf_nuevo, pdf)
    _escribir_atomico(medidas, json.dumps(bandas).encode("utf-8"))
    os.replace(docx_nuevo, docx)


def _primera_hoja(origen: Path, destino: Path) -> None:
    """Un Word membretado a veces trae dos hojas; el fondo es la primera."""
    from pypdf import PdfReader, PdfWriter

    escritor = PdfWriter()
    escritor.add_page(PdfReader(str(origen)).pages[0])
    salida = io.BytesIO()
    escritor.write(salida)
    _escribir_atomico(destino, salida.getvalue())


def borrar(empresa_id: int) -> None:
    for ruta in _rutas(empresa_id):
        ruta.unlink(missing_ok=True)


# ─── Aplicarla ──────────────────────────────────────────────

def margenes_pt(membrete: dict | None, arriba: float, abajo: float) -> tuple[float, float]:
    """Margen superior e inferior para que el texto no pise el membrete."""
    if not membrete:
        return arriba, abajo
    return (max(arriba, membrete["cabecera_pt"] + AIRE_PT),
            max(abajo, membrete["pie_pt"] + AIRE_PT))


def estampar(ruta_pdf: Path | str, membrete: dict) -> None:
    """Pone cada hoja del PDF encima de la hoja membretada, en el mismo archivo."""
    from pypdf import PdfReader, PdfWriter

    contenido = PdfReader(str(ruta_pdf))
    escritor = PdfWriter()
    for pagina in contenido.pages:
        # Se relee el fondo en cada hoja: `merge_page` modifica la pagina de
        # destino, y reusar la misma acumularia el texto de todas las hojas.
        fondo = PdfReader(str(membrete["pdf"])).pages[0]
        ancho, alto = float(pagina.mediabox.width), float(pagina.mediabox.height)
        if (abs(float(fondo.mediabox.width) - ancho) > 2
                or abs(float(fondo.mediabox.height) - alto) > 2):
            fondo.scale_to(ancho, alto)
        fondo.merge_page(pagina)
        escritor.add_page(fondo)
    salida = io.BytesIO()
    escritor.write(salida)
    # Atomico: se escribe sobre el mismo PDF que se acaba de leer, y un fallo a
    # mitad lo dejaba cortado.
    _escribir_atomico(Path(ruta_pdf), salida.getvalue())


def documento_word(empresa_id: int):
    """Un python-docx nuevo: sobre la hoja membretada si la hay, en blanco si no.

    Del Word membretado se conservan cabecera, pie y tamano de hoja, y se
    vacia el cuerpo. Los margenes laterales se suben a 2 cm como minimo: el de
    C&M San Jose trae 0 a la izquierda, pensado para la imagen de la cabecera,
    y el texto quedaria pegado al borde.
    """
    from docx import Document
    from docx.shared import Cm, Pt

    membrete = de(empresa_id)
    if not membrete:
        return Document()
    doc = Document(str(membrete["docx"]))
    cuerpo = doc.element.body
    for hijo in list(cuerpo):
        if not hijo.tag.endswith("}sectPr"):
            cuerpo.remove(hijo)
    for seccion in doc.sections:
        for lado in ("left_margin", "right_margin"):
            if (getattr(seccion, lado) or 0) < Cm(2):
                setattr(seccion, lado, Cm(2))
        # El texto empieza debajo de lo que de verdad ocupa la cabecera, y
        # termina encima del pie aunque el Word declare margen inferior 0.
        for lado, banda_pt in (("top_margin", membrete["cabecera_pt"]),
                               ("bottom_margin", membrete["pie_pt"])):
            minimo = Pt(banda_pt + AIRE_PT)
            if (getattr(seccion, lado) or 0) < minimo:
                setattr(seccion, lado, minimo)
    return doc


def miniatura_png(empresa_id: int, ancho_px: int = 420) -> bytes | None:
    """La hoja convertida, en PNG, para que el usuario vea como quedo.

    Se muestra la conversion y no el Word porque es lo que de verdad va a
    salir: si LibreOffice movio una imagen o tapo un texto, se ve aqui y no
    al imprimir la cotizacion.
    """
    import pypdfium2 as pdfium

    membrete = de(empresa_id)
    if not membrete:
        return None
    documento = pdfium.PdfDocument(str(membrete["pdf"]))
    try:
        pagina = documento[0]
        escala = ancho_px / pagina.get_width()
        imagen = pagina.render(scale=escala).to_pil()
    finally:
        documento.close()
    salida = io.BytesIO()
    imagen.save(salida, format="PNG", optimize=True)
    return salida.getvalue()
