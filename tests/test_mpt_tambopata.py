"""Municipalidad Provincial de Tambopata: la fuente, la SC del SIGA, los
anexos en Word y la postulacion armada.

Los PDF del SIGA se dibujan aqui con reportlab imitando la SC real (bienes y
servicios), con datos inventados. Los anexos se llenan sobre las plantillas
de la municipalidad que estan en el repositorio.
"""
import io
import zipfile
from datetime import date
from decimal import Decimal

import pdfplumber
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from shared import fechas
from shared.llenar_formatos import TINTA, FormatoNoReconocido
from tests.conftest import sin_base

ALTO = letter[1]


def _y(top):
    return ALTO - top


def _siga(items, servicio=False) -> bytes:
    """Una SC del SIGA: cabecera de la entidad, caja de datos del postor,
    tabla sin lineas entre items, TOTAL, condiciones y "Atentamente;"."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setFont("Helvetica-Bold", 7)
    c.drawString(32, _y(27), "Sistema Integrado de Gestión Administrativa")
    c.drawString(32, _y(37), "Módulo de Logística")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(192, _y(68), "SOLICITUD DE COTIZACIÓN Nº 000001")
    c.setFont("Helvetica-Bold", 7)
    for x, y, t in ((33, 127, "Señores"), (85, 127, ":"), (433, 127, "R.U.C."),
                    (464, 127, ":"), (34, 140, "Dirección"), (70, 140, ":"),
                    (33, 166, "Teléfono"), (85, 166, ":"), (238, 166, "Fax"), (264, 166, ":"),
                    (33, 179, "Email"), (53, 179, ":"), (238, 179, "Fecha"), (264, 179, ":")):
        c.drawString(x, _y(y), t)
    if servicio:
        xs = [22, 79, 146, 516, 574]
        cab = (("UNIDAD", 36), ("ITEM", 103), ("DESCRIPCION", 303), ("VALOR", 532))
    else:
        xs = [22, 79, 136, 196, 469, 521, 574]
        cab = (("CANTIDAD", 30), ("UNIDAD", 93), ("ITEM", 156), ("DESCRIPCION", 305),
               ("PRECIO", 480), ("PRECIO", 533))
    for x in xs:
        c.line(x, _y(205), x, _y(583))
    c.line(22, _y(205), 574, _y(205))
    c.line(22, _y(231), 574, _y(231))
    c.line(22, _y(583), 574, _y(583))
    c.line(458 if servicio else 469, _y(600), 574, _y(600))
    for t, x in cab:
        c.drawString(x, _y(216), t)
    if not servicio:
        c.drawString(477, _y(225), "UNITARIO")
        c.drawString(535, _y(225), "TOTAL")
    else:
        c.drawString(532, _y(225), "TOTAL")
    c.setFont("Helvetica", 7)
    top = 240
    for k, (cant, und, desc) in enumerate(items):
        codigo = f"71720005{k:04d}"
        if servicio:
            c.drawString(25, _y(top + 2), und)
            c.drawString(86, _y(top + 2), codigo)
            c.drawString(149, _y(top), desc)
        else:
            c.drawString(45, _y(top + 1), str(cant))
            c.drawString(82, _y(top + 2), und)
            c.drawString(139, _y(top + 2), codigo)
            c.drawString(198, _y(top), desc)
        top += 25
    c.setFont("Helvetica-Bold", 7)
    c.drawString(475 if servicio else 483, _y(594), "TOTAL")
    for y, t in ((640, "- Forma de Pago:"), (649, "- Garantía:"),
                 (658, "- Plazo de Entrega en Nº Dias/ Ejecución del Servicio :"),
                 (667, "- Tipo de Moneda :"), (676, "- Validez de la cotización :"),
                 (685, "- Indicar Marca de Procedencia"), (694, "- Tipo de Cambio :")):
        c.drawString(29, _y(y), t)
    c.drawString(295, _y(636), "Requerimientos Técnicos:")
    c.setFont("Helvetica", 7)
    c.drawString(29, _y(713), "Atentamente;")
    c.save()
    return buf.getvalue()


EMPRESA = {"razon_social": "FERRETERIA EL EJEMPLO S.A.C.", "ruc": "20123456789",
           "representante_legal": "QUISPE MAMANI ROSA ELENA", "dni_representante": "12345678",
           "telefono": "987654321", "email": "ventas@ejemplo.pe",
           "direccion": "Av. La Cultura 1234, Cusco"}
BIENES = [(60, "UNIDAD", "ARCHIVADOR DE CARTON LOMO ANCHO"),
          (30, "UNIDAD", "CORRECTOR LIQUIDO TIPO LAPICERO"),
          (50, "EMP X 500", "PAPEL BOND 80 g TAMAÑO A4")]


def _tinta(pdf: bytes) -> str:
    objetivo = tuple(round(v, 2) for v in TINTA)
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        p = doc.pages[0]
        azul = p.filter(lambda o: o.get("object_type") != "char" or tuple(
            round(v, 2) for v in (o.get("non_stroking_color") or ())) == objetivo)
        return azul.extract_text() or ""


def _items_de(sc, precio="2.50"):
    from shared.llenar_siga import leer_items_siga
    return [{**it, "marca": f"MARCA {k}", "precio_unitario": Decimal(precio),
             "subtotal": it["cantidad"] * Decimal(precio)}
            for k, it in enumerate(leer_items_siga(sc), 1)]


# ─── La SC del SIGA ─────────────────────────────────────────

def test_se_leen_los_items_de_la_sc_del_siga():
    from shared.llenar_siga import es_siga, leer_items_siga
    sc = _siga(BIENES)
    assert es_siga(sc)
    items = leer_items_siga(sc)
    assert [(i["cantidad"], i["unidad"], i["descripcion"]) for i in items] == [
        (Decimal(60), "UNIDAD", "ARCHIVADOR DE CARTON LOMO ANCHO"),
        (Decimal(30), "UNIDAD", "CORRECTOR LIQUIDO TIPO LAPICERO"),
        (Decimal(50), "EMP X 500", "PAPEL BOND 80 g TAMAÑO A4")]


def test_un_servicio_sin_cantidad_cuenta_como_uno():
    from shared.llenar_siga import leer_items_siga
    items = leer_items_siga(_siga([(None, "SERVICIO", "SERVICIO DE REFRIGERIOS")],
                                  servicio=True))
    assert [(i["cantidad"], i["unidad"], i["descripcion"]) for i in items] == [
        (Decimal(1), "SERVICIO", "SERVICIO DE REFRIGERIOS")]


def test_la_sc_del_siga_sale_llena_en_su_formato():
    from shared.llenar_siga import llenar_solicitud_siga
    sc = _siga(BIENES)
    items = _items_de(sc)
    total = sum(i["subtotal"] for i in items)
    pdf = llenar_solicitud_siga(sc, EMPRESA, items, {"forma_pago": "Depósito CCI",
                                                     "validez": "10 días"}, total)
    tinta = _tinta(pdf)
    for t in ("FERRETERIA EL EJEMPLO S.A.C.", "20123456789", "ventas@ejemplo.pe",
              "Marca: MARCA 1", "Marca: MARCA 3", "150.00", "75.00", "125.00", "350.00",
              "DEPÓSITO CCI", "10 DÍAS", "INDICADA EN CADA ÍTEM"):
        assert t in tinta, t
    # Cada precio en la fila de su item: el total de la fila 2 junto a su descripcion.
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        palabras = doc.pages[0].extract_words()
    fila2 = next(w["top"] for w in palabras if w["text"] == "CORRECTOR")
    assert any(w["text"] == "75.00" and abs(w["top"] - fila2) < 6 for w in palabras)


def test_la_sc_del_siga_no_se_llena_si_los_items_no_cuadran():
    from shared.llenar_siga import llenar_solicitud_siga
    sc = _siga(BIENES)
    items = _items_de(sc)
    with pytest.raises(FormatoNoReconocido):
        llenar_solicitud_siga(sc, EMPRESA, items[:2], {}, Decimal(1))
    items[1]["cantidad"] = Decimal(31)
    with pytest.raises(FormatoNoReconocido):
        llenar_solicitud_siga(sc, EMPRESA, items, {}, Decimal(1))


def test_el_servicio_se_cotiza_en_su_valor_total():
    from shared.llenar_siga import llenar_solicitud_siga
    sc = _siga([(None, "SERVICIO", "SERVICIO DE REFRIGERIOS")], servicio=True)
    items = [{"descripcion": "SERVICIO DE REFRIGERIOS", "cantidad": Decimal(1),
              "precio_unitario": Decimal(1500), "subtotal": Decimal(1500)}]
    assert "1,500.00" in _tinta(llenar_solicitud_siga(sc, EMPRESA, items, {}, Decimal(1500)))


# ─── La fuente ──────────────────────────────────────────────

CONVOCATORIA = {
    "id": "2c5d0757-c5d7-4ff9-9c0a-dc7c7700ddfb", "type": "BIENES", "estado": "AMPLIADA",
    "estudioMercado": "Convocatoria 1 – E/M N.º 4201-2026",
    "descripcion": "Convocatoria 1 – E/M N.º 4201-2026-ADQUISICION DE PAPEL BOND 80g "
                   "TAMAÑO A4(OFICINA DE IMAGEN INSTITUCIONAL)",
    "fechaInicio": "2026-10-15T15:16:00.000Z", "fechaFin": "2026-07-07T21:30:00.000Z",
    "fechaAmpliacion": "2026-10-16T21:30:00.000Z",
    "tdr": "/storage/2025/abastecimiento/1-tdr.pdf",
    "expediente": "/storage/2025/abastecimiento/1-exp.pdf",
}


def test_una_convocatoria_abierta_se_vuelve_licitacion():
    from radar_bot.scrapers.mpt_tambopata import a_licitacion, numero_em
    lic = a_licitacion(CONVOCATORIA)
    assert lic["nomenclatura"] == "E/M-4201-2026-MPT"
    assert lic["objeto"].startswith("[BIENES] ADQUISICION DE PAPEL BOND")
    assert lic["departamento"] == "Madre de Dios" and lic["tipo"] == "cotizacion"
    # El cierre es la ampliacion, en hora de Lima (21:30 UTC = 16:30).
    assert lic["fecha_cierre"] == fechas.fija(2026, 10, 16, 16, 30)
    assert [d["tipo"] for d in lic["_documentos"]] == ["sc", "eett"]
    assert lic["_documentos"][0]["url"].endswith("/storage/2025/abastecimiento/1-exp.pdf")
    assert numero_em(lic["nomenclatura"]) == "E/M N.º 4201-2026"
    assert a_licitacion({**CONVOCATORIA, "estado": "CERRADA"}) is None


def test_los_items_de_la_sc_quedan_en_el_formato_que_siembra_la_propuesta():
    from radar_bot.scrapers.mpt_tambopata import descripcion_desde_sc
    from shared.cotizacion import items_de_texto
    texto = descripcion_desde_sc(_siga(BIENES + [(5, "UNIDAD", "TIJERA (METAL) DE 10 in")]))
    items = items_de_texto(texto)
    assert [(i["descripcion"], i["cantidad"], i["unidad"]) for i in items][-2:] == [
        ("PAPEL BOND 80 g TAMAÑO A4", Decimal(50), "EMP X 500"),
        ("TIJERA [METAL] DE 10 in", Decimal(5), "UNIDAD")]


# ─── Los anexos ─────────────────────────────────────────────

def _datos(**extra):
    return {"empresa": EMPRESA, "items": [
                {"descripcion": "PAPEL BOND A4", "cantidad": Decimal(50), "unidad": "EMP X 500",
                 "marca": "Atlas", "precio_unitario": Decimal("18.50"),
                 "subtotal": Decimal("925.00")}],
            "condiciones": {"plazo_entrega": "05 días calendario", "garantia": "06 meses"},
            "total": Decimal("925.00"), "referencia": "E/M N.º 4201-2026",
            "cci": "00212345678901234567", "banco": "BCP", "partida": "11022172",
            "hoy": date(2026, 10, 15), "imagenes": {}, **extra}


@pytest.fixture
def plantillas_del_repo(monkeypatch, tmp_path):
    from shared import anexos_mpt
    monkeypatch.setattr(anexos_mpt, "CACHE", tmp_path / "sin_cache")
    monkeypatch.setattr(anexos_mpt, "_refrescar_cache", lambda: None)


def _texto(docx_bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(docx_bytes))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            partes += [c.text for c in r.cells]
    return "\n".join(partes)


def test_los_anexos_salen_llenos_con_los_datos_del_postor(plantillas_del_repo):
    from shared.anexos_mpt import llenar_docx
    t05 = _texto(llenar_docx("05", _datos()))
    for v in ("Referencia: E/M N.º 4201-2026", "PAPEL BOND A4 - Marca: Atlas", "50 EMP X 500",
              "18.50", "925.00", "TOTAL S/", "05 días calendario", "06 meses",
              "FERRETERIA EL EJEMPLO S.A.C.", "20123456789", "ventas@ejemplo.pe",
              "987654321", "15/10/2026"):
        assert v in t05, v
    assert "[INDICAR" not in t05
    t06 = _texto(llenar_docx("06", _datos()))
    assert "El que suscribe, QUISPE MAMANI ROSA ELENA" in t06
    assert "identificado con DNI N° 12345678" in t06
    assert "Partida Registral Nº 11022172" in t06 and "[CONSIGNAR" not in t06
    t07 = _texto(llenar_docx("07", _datos()))
    assert "00212345678901234567" in t07 and "BCP" in t07
    assert "declarante QUISPE MAMANI ROSA ELENA, con DNI 12345678" in _texto(
        llenar_docx("08", _datos()))
    assert "Representante Legal de FERRETERIA EL EJEMPLO S.A.C., identificado/a con DNI N° " \
           "12345678" in _texto(llenar_docx("09", _datos()))


def test_con_sello_la_firma_se_estampa_encima_sin_mover_la_hoja(plantillas_del_repo, tmp_path):
    import docx
    from PIL import Image

    from shared.anexos_mpt import llenar_docx
    sello = tmp_path / "sello.png"
    Image.new("RGB", (300, 120), "blue").save(sello)
    d = docx.Document(io.BytesIO(llenar_docx("08", _datos(imagenes={"sello": str(sello)}))))
    xml = d.element.xml
    assert "<wp:anchor" in xml and "<wp:wrapNone/>" in xml and "<wp:inline" not in xml


@pytest.mark.asyncio
async def test_sin_libreoffice_los_anexos_van_en_word_y_se_avisa(plantillas_del_repo,
                                                                 monkeypatch):
    from shared import anexos_mpt, pdf_firma
    monkeypatch.setattr(pdf_firma, "hay_conversor", lambda: None)
    salida, aviso = await anexos_mpt.anexos_pdf(_datos())
    assert sorted(salida) == [f"{n}.docx" for n in sorted(anexos_mpt.ANEXOS.values())]
    assert aviso and "LibreOffice" in aviso


# ─── La postulacion armada, en la web ───────────────────────

@sin_base
@pytest.mark.asyncio
async def test_la_postulacion_a_la_mpt_sale_completa_y_el_panel_lleva_al_portal(
        usuario, empresa, cliente, marca, tmp_path, monkeypatch, plantillas_del_repo):
    from radar_bot.scrapers.mpt_tambopata import descripcion_desde_sc
    from shared import constancia_rnp, cotizacion, ficha_ruc, llenar_formatos, pdf_firma
    from shared.db import connection, guardar_documento

    monkeypatch.setattr(llenar_formatos, "DIR", tmp_path / "formatos")
    monkeypatch.setattr(constancia_rnp, "DIR_PDF", tmp_path / "rnp")
    monkeypatch.setattr(ficha_ruc, "DIR_PDF", tmp_path / "ficha")
    monkeypatch.setattr(pdf_firma, "hay_conversor", lambda: None)
    sc = _siga(BIENES)
    lic = f"mpt_prueba_{marca}"
    async with connection() as c:
        await c.execute(
            """INSERT INTO licitaciones (id, fuente, tipo, nomenclatura, entidad, objeto,
                   descripcion, url)
               VALUES ($1, 'mpt_tambopata', 'cotizacion', 'E/M-4201-2026-MPT',
                   'Municipalidad Provincial de Tambopata', '[BIENES] UTILES', $2, $3)""",
            lic, descripcion_desde_sc(sc),
            "https://portal.munitambopata.gob.pe/servicios/contratacion-bienes-servicios")
        await c.execute("DELETE FROM intentos_acceso")
        pid = await c.fetchval("INSERT INTO propuestas (licitacion_id, empresa_id, estado) "
                               "VALUES ($1, $2, 'iniciado') RETURNING id", lic, empresa)
    try:
        await guardar_documento(lic, "sc", None, "expediente.pdf",
                                "https://portal.munitambopata.gob.pe/storage/x-exp.pdf", sc)
        await cotizacion.sembrar(pid)
        items = await cotizacion.items_de(pid)
        assert len(items) == 3
        await cotizacion.guardar(pid, {it["id"]: (f"MARCA {k}", Decimal("2.50"))
                                       for k, it in enumerate(items, 1)}, {})
        r = await cliente.post("/entrar", data={"email": usuario["email"],
                                                "password": usuario["password"]})
        assert r.status_code == 303

        pagina = (await cliente.get(f"/propuestas/{pid}")).text
        assert "Abrir el portal de la Municipalidad" in pagina
        assert 'data-copiar="E/M N.º 4201-2026"' in pagina
        assert "servicios/contratacion-bienes-servicios" in pagina
        assert "Enviar propuesta" in pagina and "No soy un robot" not in pagina

        r = await cliente.get(f"/propuestas/{pid}/paquete")
        assert r.status_code == 200, r.headers.get("location")
        z = zipfile.ZipFile(io.BytesIO(r.content))
        nombres = z.namelist()
        assert nombres[0] == "01_Solicitud_de_Cotizacion.pdf"
        assert [n[3:] for n in nombres[1:6]] == [
            "Anexo_05_Oferta_Economica.docx", "Anexo_06_DJ_Notificacion.docx",
            "Anexo_07_Carta_Autorizacion_CCI.docx", "Anexo_08_DJ_Contratos_Menores.docx",
            "Anexo_09_DJ_Prohibiciones.docx"]
        assert "Marca: MARCA 2" in _tinta(z.read(nombres[0]))
        assert "E/M N.º 4201-2026" in _texto(z.read(nombres[1]))
    finally:
        async with connection() as c:
            await c.execute("DELETE FROM propuestas WHERE id=$1", pid)
            await c.execute("DELETE FROM licitaciones WHERE id=$1", lic)


def test_el_paquete_armado_a_mano_se_sirve_tal_cual(tmp_path, monkeypatch):
    """Si el usuario firmo a mano, su ZIP manda: la descarga no genera otro."""
    from web import propuestas
    monkeypatch.setattr(propuestas, "PAQUETES", tmp_path)
    assert propuestas.paquete_armado(7) is None
    with zipfile.ZipFile(tmp_path / "propuesta_7.zip", "w") as z:
        z.writestr("02_Anexo_05.pdf", b"x")
        z.writestr("01_Solicitud.pdf", b"x")
    ruta = propuestas.paquete_armado(7)
    assert ruta == tmp_path / "propuesta_7.zip"
    assert propuestas._contenido_zip(ruta) == ["01_Solicitud.pdf", "02_Anexo_05.pdf"]
