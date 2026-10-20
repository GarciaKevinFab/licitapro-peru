"""Auditoria: los llenadores no entregan documentos con huecos en silencio.

Antes, si faltaba el DNI, el CCI o el domicilio, el anexo o la DJ salian
igual, con "……", "____" o la linea en blanco (el Anexo 07 de CUBS FAM se
envio con puntos en el CCI). Si la municipalidad cambiaba un marcador, el
"[CONSIGNAR ...]" quedaba en el documento firmado. Y un formato con una forma
rara terminaba en un 500 en vez de caer al formato propio.

Aqui se prueba que:
  * falta un dato obligatorio      -> DocumentoIncompleto, con el dato
  * la plantilla ya no lo tiene    -> PlantillaCambiada (FormatoNoReconocido)
  * el PDF tiene una forma rara    -> FormatoNoReconocido, nunca ValueError
"""
import io
import threading
import zipfile
from datetime import date
from decimal import Decimal

import pdfplumber
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from shared.incompleto import DocumentoIncompleto
from shared.llenar_formatos import FormatoNoReconocido
from tests.test_llenar_formatos import EMPRESA, _dj, _items, _sc
from tests.test_mpt_tambopata import _datos, _siga


@pytest.fixture
def plantillas_del_repo(monkeypatch, tmp_path):
    from shared import anexos_mpt
    monkeypatch.setattr(anexos_mpt, "CACHE", tmp_path / "sin_cache")
    monkeypatch.setattr(anexos_mpt, "_refrescar_cache", lambda: None)


def _sin(dic: dict, *claves) -> dict:
    return {k: v for k, v in dic.items() if k not in claves}


# ─── Anexos de la MPT: datos obligatorios ───────────────────

def test_sin_dni_el_anexo_no_se_genera_y_se_dice_que_falta(plantillas_del_repo):
    from shared.anexos_mpt import llenar_docx
    datos = _datos()
    datos["empresa"] = _sin(datos["empresa"], "dni_representante")
    for numero in ("05", "06", "07", "08", "09"):
        with pytest.raises(DocumentoIncompleto) as e:
            llenar_docx(numero, datos)
        assert e.value.documento == f"el Anexo {numero}"
        assert "el DNI del representante" in e.value.faltantes
        assert "DNI" in e.value.mensaje()


def test_el_anexo_07_exige_un_cci_de_20_digitos_y_el_banco(plantillas_del_repo):
    from shared.anexos_mpt import llenar_docx
    for cci in (None, "", "0021234567890123456", "002-123-456789012345-6X"):
        with pytest.raises(DocumentoIncompleto) as e:
            llenar_docx("07", _datos(cci=cci))
        assert e.value.faltantes == ["el CCI (20 dígitos)"]
    with pytest.raises(DocumentoIncompleto) as e:
        llenar_docx("07", _datos(banco=None))
    assert e.value.faltantes == ["el banco del CCI"]
    # Con espacios o guiones se acepta y se escribe limpio.
    import docx
    d = docx.Document(io.BytesIO(llenar_docx("07", _datos(cci="002 1234 5678 9012 3456 7"))))
    celdas = [c.text for f in d.tables[0].rows for c in f.cells]
    assert "00212345678901234567" in celdas and "BCP" in celdas


def test_el_anexo_05_exige_precios_y_plazo(plantillas_del_repo):
    from shared.anexos_mpt import llenar_docx
    datos = _datos()
    datos["items"] = [{**datos["items"][0], "precio_unitario": None, "subtotal": None}]
    datos["condiciones"] = {"plazo_entrega": "", "garantia": ""}
    with pytest.raises(DocumentoIncompleto) as e:
        llenar_docx("05", datos)
    assert "el precio de «PAPEL BOND A4»" in e.value.faltantes
    assert "el plazo de entrega" in e.value.faltantes


def test_lo_opcional_puede_faltar(plantillas_del_repo):
    """Garantia, telefono, correo y datos registrales no invalidan la oferta."""
    from shared.anexos_mpt import PUNTOS, llenar_docx
    datos = _datos(partida=None)
    datos["empresa"] = _sin(datos["empresa"], "telefono", "email")
    datos["condiciones"] = {"plazo_entrega": "05 días calendario"}
    import docx
    d = docx.Document(io.BytesIO(llenar_docx("05", datos)))
    garantia = next(f for f in d.tables[1].rows if "Garant" in f.cells[0].text)
    assert garantia.cells[-1].text == PUNTOS
    texto06 = "\n".join(p.text for p in docx.Document(
        io.BytesIO(llenar_docx("06", datos))).paragraphs)
    assert f"Partida Registral Nº {PUNTOS}" in texto06


@pytest.mark.asyncio
async def test_anexos_pdf_dice_todo_lo_que_falta_de_una_vez(plantillas_del_repo):
    from shared.anexos_mpt import anexos_pdf
    datos = _datos(cci=None)
    datos["empresa"] = _sin(datos["empresa"], "direccion")
    with pytest.raises(DocumentoIncompleto) as e:
        await anexos_pdf(datos)
    assert e.value.documento == "los Anexos 05, 06, 07"
    assert e.value.faltantes == ["el domicilio fiscal", "el CCI (20 dígitos)"]


# ─── Anexos de la MPT: la plantilla cambio ──────────────────

def _plantilla_cambiada(tmp_path, monkeypatch, numero, cambiar):
    """Copia la plantilla del repositorio, le aplica `cambiar(doc)` y hace que
    el llenador la use."""
    import docx

    from shared import anexos_mpt
    doc = docx.Document(str(anexos_mpt.REPO / f"anexo_{numero}.docx"))
    cambiar(doc)
    ruta = tmp_path / f"anexo_{numero}_cambiado.docx"
    doc.save(ruta)
    monkeypatch.setattr(anexos_mpt, "plantilla", lambda n: ruta)


def test_si_el_marcador_cambio_de_redaccion_no_se_entrega_con_puntos(
        plantillas_del_repo, tmp_path, monkeypatch):
    from shared.anexos_mpt import PlantillaCambiada, llenar_docx

    def otra_redaccion(doc):
        p = next(p for p in doc.paragraphs if "declarante" in p.text)
        for r in p.runs:
            r.text = r.text.replace("declarante", "suscrito")
    _plantilla_cambiada(tmp_path, monkeypatch, "08", otra_redaccion)
    with pytest.raises(PlantillaCambiada, match="Anexo 08") as e:
        llenar_docx("08", _datos())
    # Es un FormatoNoReconocido: la web lo muestra con su mensaje, no da un 500.
    assert isinstance(e.value, FormatoNoReconocido)


def test_un_marcador_nuevo_en_la_plantilla_no_llega_al_documento(
        plantillas_del_repo, tmp_path, monkeypatch):
    from shared.anexos_mpt import PlantillaCambiada, llenar_docx
    _plantilla_cambiada(tmp_path, monkeypatch, "09",
                        lambda doc: doc.add_paragraph("Fecha: [CONSIGNAR LA FECHA DE FIRMA]"))
    with pytest.raises(PlantillaCambiada, match=r"\[CONSIGNAR LA FECHA DE FIRMA\]"):
        llenar_docx("09", _datos())


def test_el_anexo_07_va_por_etiqueta_y_no_por_posicion(
        plantillas_del_repo, tmp_path, monkeypatch):
    """Una fila de mas al inicio ya no corre el CCI a la fila del banco; una
    fila que desaparece es un error y no un dato perdido."""
    import copy

    import docx

    from shared.anexos_mpt import PlantillaCambiada, llenar_docx

    def fila_de_mas(doc):
        t = doc.tables[0]
        nueva = copy.deepcopy(t.rows[0]._tr)
        t.rows[0]._tr.addprevious(nueva)
        for p in t.rows[0].cells[0].paragraphs:
            for r in p.runs:
                r.text = ""
        t.rows[0].cells[0].paragraphs[0].add_run("Fecha de la carta")
    _plantilla_cambiada(tmp_path, monkeypatch, "07", fila_de_mas)
    d = docx.Document(io.BytesIO(llenar_docx("07", _datos())))
    filas = {f.cells[0].text.split()[0]: f.cells[-1].text for f in d.tables[0].rows}
    assert filas["Fecha"] == ""
    assert filas["Entidad"] == "BCP"
    assert any(v == "00212345678901234567" for v in filas.values())

    def sin_banco(doc):
        t = doc.tables[0]
        fila = next(f for f in t.rows if "Bancaria" in f.cells[0].text)
        t._tbl.remove(fila._tr)
    _plantilla_cambiada(tmp_path, monkeypatch, "07", sin_banco)
    with pytest.raises(PlantillaCambiada, match="Entidad Bancaria"):
        llenar_docx("07", _datos())


def test_el_anexo_10_suma_en_decimal_solo_lo_que_declara(plantillas_del_repo):
    """Siete experiencias y seis filas: el total es el de las seis que se ven."""
    import docx

    from shared.anexos_mpt import llenar_docx
    exp = [{"entidad_contratante": f"ENTIDAD {k}", "objeto_contrato": "X",
            "numero_contrato": f"O/C {k}", "monto": Decimal("0.10"),
            "fecha_inicio": date(2026, 1, 1), "fecha_fin": date(2026, 1, 2),
            "conformidad": True} for k in range(7)]
    exp[0]["monto"] = 0.2   # de la base puede venir float
    d = docx.Document(io.BytesIO(llenar_docx("10", _datos(experiencia=exp))))
    assert d.tables[0].rows[-1].cells[-1].text == "0.70"
    assert "ENTIDAD 6" not in "\n".join(c.text for f in d.tables[0].rows for c in f.cells)


def test_los_temporales_de_la_firma_se_borran(plantillas_del_repo, tmp_path, monkeypatch):
    import tempfile

    from PIL import Image

    from shared.anexos_mpt import llenar_docx
    firma, sello = tmp_path / "firma.png", tmp_path / "sello.png"
    Image.new("RGB", (200, 80), "black").save(firma)
    Image.new("RGB", (300, 120), "blue").save(sello)
    temporales = tmp_path / "tmp"
    temporales.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(temporales))
    llenar_docx("08", _datos(imagenes={"firma": str(firma), "sello": str(sello)}))
    assert list(temporales.iterdir()) == []
    # Y las imagenes quedan cerradas: en Windows se pueden borrar al momento.
    firma.unlink()
    sello.unlink()


@pytest.mark.asyncio
async def test_todos_los_avisos_de_conversion_llegan(plantillas_del_repo, monkeypatch):
    from shared import anexos_mpt, pdf_firma

    async def falla(ruta, carpeta):
        return None, "LibreOffice tardó demasiado"
    monkeypatch.setattr(pdf_firma, "hay_conversor", lambda: "soffice")
    monkeypatch.setattr(pdf_firma, "docx_a_pdf", falla)
    salida, aviso = await anexos_mpt.anexos_pdf(_datos())
    assert len(salida) == len(anexos_mpt.ANEXOS)
    for numero in anexos_mpt.ANEXOS:
        assert f"El Anexo {numero} va en Word" in aviso


# ─── Anexos de la MPT: la cache del portal ──────────────────

def _zip(archivos: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for nombre, contenido in archivos.items():
            z.writestr(nombre, contenido)
    return buf.getvalue()


class _Respuesta:
    def __init__(self, contenido):
        self.content = contenido

    def raise_for_status(self):
        pass


def test_un_zip_sin_anexos_no_marca_la_cache_y_no_se_reintenta_enseguida(
        tmp_path, monkeypatch):
    import httpx

    from shared import anexos_mpt
    monkeypatch.setattr(anexos_mpt, "CACHE", tmp_path / "cache")
    monkeypatch.setattr(anexos_mpt, "_ultimo_fallo", 0.0)
    pedidos = []

    def get(url, **kw):
        pedidos.append(url)
        return _Respuesta(_zip({"Leeme.txt": b"hola"}))
    monkeypatch.setattr(httpx, "get", get)
    anexos_mpt._refrescar_cache()
    assert not (tmp_path / "cache" / ".bajado").exists()
    anexos_mpt._refrescar_cache()          # el portal fallo hace un momento
    assert len(pedidos) == 1


def test_el_zip_bueno_deja_las_plantillas_enteras(tmp_path, monkeypatch):
    import httpx

    from shared import anexos_mpt
    monkeypatch.setattr(anexos_mpt, "CACHE", tmp_path / "cache")
    monkeypatch.setattr(anexos_mpt, "_ultimo_fallo", 0.0)
    plantilla = (anexos_mpt.REPO / "anexo_07.docx").read_bytes()
    monkeypatch.setattr(httpx, "get", lambda url, **kw: _Respuesta(
        _zip({"formatos/Anexo 07 Carta CCI.docx": plantilla})))
    anexos_mpt._refrescar_cache()
    assert (tmp_path / "cache" / ".bajado").exists()
    assert (tmp_path / "cache" / "anexo_07.docx").read_bytes() == plantilla
    assert [p.name for p in (tmp_path / "cache").iterdir() if p.suffix == ".tmp"] == []


# ─── SC y DJ de la entidad ──────────────────────────────────

def test_la_dj_sin_dni_no_se_genera():
    from shared.llenar_formatos import llenar_declaracion_jurada
    with pytest.raises(DocumentoIncompleto) as e:
        llenar_declaracion_jurada(_dj(), _sin(EMPRESA, "dni_representante", "direccion"),
                                  None, "10", "10")
    assert e.value.documento == "la declaración jurada"
    assert e.value.faltantes == ["el DNI del representante", "el domicilio fiscal"]


def test_la_sc_sin_domicilio_no_se_genera():
    from shared.llenar_formatos import llenar_solicitud
    items = _items()
    with pytest.raises(DocumentoIncompleto, match="domicilio"):
        llenar_solicitud(_sc(), _sin(EMPRESA, "direccion"), items, {}, "X", Decimal(1))


def test_la_sc_del_siga_sin_representante_no_se_genera():
    from shared.llenar_siga import llenar_solicitud_siga
    sc = _siga([(60, "UNIDAD", "ARCHIVADOR")])
    items = [{"descripcion": "ARCHIVADOR", "cantidad": Decimal(60),
              "precio_unitario": Decimal(1), "subtotal": Decimal(60)}]
    with pytest.raises(DocumentoIncompleto, match="representante legal"):
        llenar_solicitud_siga(sc, _sin(EMPRESA, "representante_legal"), items, {}, Decimal(60))


def _sc_sin_direccion(monkeypatch) -> bytes:
    """La SC del GOREMAD (firma junto al Area de Logistica, sin bloque de
    datos del postor) sin la etiqueta "Dirección" en la cabecera: no hay
    ningun sitio donde poner el domicilio."""
    from tests.test_llenar_formatos import _sc_con_firma_de_la_entidad
    original = canvas.Canvas.drawString

    def dibujar(self, x, y, texto, *a, **k):
        if texto != "Dirección":
            return original(self, x, y, texto, *a, **k)
    with monkeypatch.context() as m:
        m.setattr(canvas.Canvas, "drawString", dibujar)
        return _sc_con_firma_de_la_entidad()


def test_si_la_sc_no_tiene_donde_poner_el_domicilio_cae_al_formato_propio(monkeypatch):
    from shared.llenar_formatos import llenar_solicitud
    sc = _sc_sin_direccion(monkeypatch)
    with pdfplumber.open(io.BytesIO(sc)) as doc:
        assert "Dirección" not in doc.pages[0].extract_text()
    items = _items()
    with pytest.raises(FormatoNoReconocido, match="domicilio fiscal"):
        llenar_solicitud(sc, EMPRESA, items, {}, "X", sum(i["subtotal"] for i in items))


def test_las_condiciones_sin_dato_quedan_en_blanco_pero_en_el_log(caplog):
    from shared.llenar_formatos import llenar_solicitud
    items = _items()
    with caplog.at_level("WARNING", logger="shared.llenar_formatos"):
        llenar_solicitud(_sc(), EMPRESA, items, {"validez": "30 días"}, "X",
                         sum(i["subtotal"] for i in items))
    assert "FORMA DE PAGO" in caplog.text


def _dj_con_partida() -> bytes:
    """Una DJ donde, tras mencionar el RUC, viene una "Partida N° ....." y la
    linea del RUC va despues."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica", 10)
    alto = A4[1]
    c.drawString(56, alto - 255, "Yo, " + "." * 60 + " representante legal de la Empresa")
    c.drawString(56, alto - 269, "." * 60 + ", identificado con DNI N° " + "." * 30 + ",")
    c.drawString(56, alto - 283, "inscrita con su RUC en la Partida N° " + "." * 30)
    c.drawString(56, alto - 297, "RUC N° " + "." * 30 + " con domicilio fiscal en " + "." * 50)
    c.save()
    return buf.getvalue()


def test_el_ruc_de_la_dj_va_en_su_linea_y_no_en_la_partida():
    from shared.llenar_formatos import TINTA, llenar_declaracion_jurada
    pdf = llenar_declaracion_jurada(_dj_con_partida(), EMPRESA, None, "10", "10")
    objetivo = tuple(round(v, 2) for v in TINTA)
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        azul = doc.pages[0].filter(lambda o: o.get("object_type") == "char" and tuple(
            round(v, 2) for v in (o.get("non_stroking_color") or ())) == objetivo)
        renglones = {}
        for ch in azul.chars:
            renglones.setdefault(round(ch["top"]), []).append(ch)
    texto_por_renglon = {top: "".join(c["text"] for c in sorted(cs, key=lambda c: c["x0"]))
                         for top, cs in renglones.items()}
    con_ruc = [top for top, t in texto_por_renglon.items() if "20123456789" in t]
    assert len(con_ruc) == 1
    assert abs(con_ruc[0] - (297 - 10)) < 8              # en el renglon "RUC N°"
    assert not any(abs(top - (283 - 10)) < 5 for top in texto_por_renglon)  # partida: nada


def test_una_sc_con_una_forma_rara_es_formato_no_reconocido():
    """MARCA y PRECIO estan, pero no hay linea vertical a la izquierda de MARCA:
    `max()` de una lista vacia era un ValueError y un 500."""
    from shared.llenar_formatos import _localizar_tabla, llenar_solicitud
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    alto = A4[1]
    for x in (374, 436, 498, 561):
        c.line(x, alto - 323, x, alto - 450)
    c.drawString(388, alto - 343, "MARCA")
    c.drawString(450, alto - 338, "PRECIO")
    c.drawString(512, alto - 338, "PRECIO")
    c.save()
    with pdfplumber.open(io.BytesIO(buf.getvalue())) as doc, pytest.raises(FormatoNoReconocido):
        _localizar_tabla(doc.pages[0])
    with pytest.raises(FormatoNoReconocido):
        llenar_solicitud(buf.getvalue(), EMPRESA, _items(), {}, "X", Decimal(1))


def test_un_item_sin_cantidad_es_formato_no_reconocido_y_no_un_500():
    from shared.llenar_formatos import llenar_solicitud
    items = _items()
    items[1]["cantidad"] = None
    with pytest.raises(FormatoNoReconocido, match="no tiene cantidad"):
        llenar_solicitud(_sc(), EMPRESA, items, {}, "X", Decimal(1))


def test_subir_un_pdf_danado_es_formato_no_reconocido(tmp_path, monkeypatch):
    from shared import llenar_formatos
    monkeypatch.setattr(llenar_formatos, "DIR", tmp_path)
    with pytest.raises(FormatoNoReconocido):
        llenar_formatos.guardar(1, "sc", b"%PDF-1.4\n" + b"\x00basura" * 50)
    assert not llenar_formatos.guardados(1)["sc"]


def test_la_cantidad_ilegible_de_un_bien_no_se_vuelve_uno():
    from shared.llenar_siga import leer_items_siga, llenar_solicitud_siga
    sc = _siga([(60, "UNIDAD", "ARCHIVADOR"), ("??", "UNIDAD", "CORRECTOR")])
    with pytest.raises(FormatoNoReconocido, match="cantidad del ítem 2"):
        leer_items_siga(sc)
    items = [{"descripcion": d, "cantidad": Decimal(c), "precio_unitario": Decimal(1),
              "subtotal": Decimal(c)} for d, c in (("ARCHIVADOR", 60), ("CORRECTOR", 1))]
    with pytest.raises(FormatoNoReconocido, match="cantidad del ítem 2"):
        llenar_solicitud_siga(sc, EMPRESA, items, {}, Decimal(61))


# ─── Fuera del hilo del servidor ────────────────────────────

@pytest.mark.asyncio
async def test_la_sc_la_dj_y_las_eett_se_llenan_en_otro_hilo(monkeypatch):
    from shared import llenar_formatos
    hilos = []

    async def formato(pid, lic, tipo):
        return {"sc": _sc(), "dj": _dj(), "eett": _sc()}[tipo], "subido"
    monkeypatch.setattr(llenar_formatos, "formato", formato)
    for nombre in ("llenar_solicitud", "llenar_declaracion_jurada", "sellar_paginas"):
        original = getattr(llenar_formatos, nombre)

        def espia(*a, _original=original, **k):
            hilos.append(threading.get_ident())
            return _original(*a, **k)
        monkeypatch.setattr(llenar_formatos, nombre, espia)
    items = _items()
    await llenar_formatos.sc_llena(1, "x", EMPRESA, items, {})
    await llenar_formatos.dj_llena(1, "x", EMPRESA, None)
    await llenar_formatos.eett_selladas(1, "x", {})
    assert len(hilos) == 3 and threading.get_ident() not in hilos


# ─── Los PDF propios (pdf_firmable, documentos_menores) ─────

@pytest.fixture
def _pdf_propio(tmp_path, monkeypatch):
    import shared.pdf_firmable as pf

    async def _sin_imagenes(_):
        return {}
    monkeypatch.setattr(pf, "rutas_de", _sin_imagenes)
    monkeypatch.setattr(pf, "SALIDA", tmp_path)
    return pf


def _texto_pdf(ruta) -> str:
    with pdfplumber.open(ruta) as doc:
        return "\n".join(p.extract_text() or "" for p in doc.pages)


def test_el_lugar_no_se_inventa_y_va_escapado():
    from shared.pdf_firmable import lugar_y_fecha
    hoy = date(2026, 10, 20)
    assert lugar_y_fecha({}, hoy) == "20 de octubre de 2026"
    assert "Lima" not in lugar_y_fecha({"departamento": ""}, hoy)
    assert lugar_y_fecha({"departamento": "Madre de Dios"}, hoy).startswith("Madre de Dios, ")
    assert lugar_y_fecha({"departamento": "Madre de Dios", "provincia": "Tambopata"},
                         hoy).startswith("Tambopata, ")
    assert lugar_y_fecha({"departamento": "A & <B>"}, hoy).startswith("A &amp; &lt;B&gt;, ")


@pytest.mark.asyncio
async def test_un_departamento_con_ampersand_no_rompe_el_pdf(_pdf_propio):
    emp = {"id": 9, **EMPRESA, "departamento": "Madre de Dios & <Tambopata>"}
    ruta = await _pdf_propio.generar_pdf("lugar.pdf", "COTIZACIÓN", "", ["x"], emp)
    texto = _texto_pdf(ruta)
    assert "Madre de Dios & <Tambopata>," in texto and "Lima" not in texto


@pytest.mark.asyncio
async def test_la_dj_propia_sin_dni_no_se_genera(_pdf_propio):
    from shared.documentos_menores import generar_dj
    prop = {"id": 5, "nomenclatura": "COT-1-2026-GOREMAD", "objeto": "X",
            "entidad": "Gobierno Regional de Madre de Dios"}
    with pytest.raises(DocumentoIncompleto) as e:
        await generar_dj(prop, {"id": 9, **_sin(EMPRESA, "dni_representante")}, None, True)
    assert e.value.faltantes == ["el DNI del representante"]
    # Con todo, el CCI puede faltar: su linea queda para escribirla a mano.
    ruta = await generar_dj(prop, {"id": 9, **EMPRESA}, None, True)
    texto = _texto_pdf(ruta)
    assert "12345678" in texto and "____" in texto


@pytest.mark.asyncio
async def test_la_cotizacion_propia_sin_domicilio_no_se_genera(_pdf_propio):
    from shared.documentos_menores import generar_cotizacion
    prop = {"id": 6, "nomenclatura": "COT-1-2026-X", "objeto": "X", "entidad": "E"}
    with pytest.raises(DocumentoIncompleto, match="domicilio"):
        await generar_cotizacion(prop, {"id": 9, **_sin(EMPRESA, "direccion")}, _items(), {},
                                 None, True)


@pytest.mark.asyncio
async def test_el_pdf_propio_se_arma_en_otro_hilo(_pdf_propio, monkeypatch):
    hilos = []
    original = _pdf_propio.SimpleDocTemplate.build

    def build(self, *a, **k):
        hilos.append(threading.get_ident())
        return original(self, *a, **k)
    monkeypatch.setattr(_pdf_propio.SimpleDocTemplate, "build", build)
    await _pdf_propio.generar_pdf("hilo.pdf", "T", "", ["x"], {"id": 9, **EMPRESA})
    assert hilos and threading.get_ident() not in hilos
