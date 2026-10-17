"""La hoja membretada de cada empresa en los documentos que ella firma.

QUE PROTEGE

  El 2026-10-17 el usuario paso las hojas membretadas de sus cuatro empresas.
  La cotizacion y la declaracion jurada propias tienen que salir impresas
  sobre esa hoja, con el texto en el hueco entre la cabecera y el pie, sin
  pisarlos y sin repetir debajo el logo que ya trae la hoja.

  Las bandas se miden sobre la hoja convertida: los Word membretados declaran
  margenes que no coinciden con lo que dibujan (C&M San Jose trae margen
  izquierdo 0; Soluciones Informaticas, margen inferior 0 con un pie de tres
  lineas).

POR QUE SIN LIBREOFFICE

  La conversion del Word la hace LibreOffice en el servidor. Aqui se simula
  con una hoja hecha con reportlab que tiene lo que importa medir: una
  cabecera, un pie y una marca de agua en el centro, que NO es membrete.
"""
import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Cm
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from shared import membrete
from tests.conftest import sin_base

ALTO = A4[1]


def _hoja_pdf(ruta: Path, cabecera=110, pie=50, marca_de_agua=True) -> None:
    """Hoja A4: un bloque oscuro de `cabecera` pt arriba y de `pie` pt abajo."""
    c = canvas.Canvas(str(ruta), pagesize=A4)
    if cabecera:
        c.setFillGray(0)
        c.rect(40, ALTO - cabecera, 500, cabecera - 10, fill=1, stroke=0)
        c.setFillGray(1)
        c.drawString(60, ALTO - 50, "MEMBRETE DE PRUEBA SAC  RUC 20000000001")
    if marca_de_agua:
        # Gris oscuro a proposito y en el centro: lo que la aleja del membrete
        # es la distancia al borde, no lo claro que sea.
        c.setFillGray(0.5)
        c.rect(100, ALTO / 2 - 40, 400, 80, fill=1, stroke=0)
    if pie:
        c.setFillGray(0)
        c.rect(40, 0, 500, pie, fill=1, stroke=0)
    c.save()


def _hoja_docx(cabecera_texto="MEMBRETE DE PRUEBA SAC") -> bytes:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = cabecera_texto
    doc.sections[0].left_margin = 0          # como el Word de C&M San Jose
    doc.add_paragraph("Esto es el cuerpo de la hoja y no debe quedar")
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


@pytest.fixture
def _dir(tmp_path, monkeypatch):
    monkeypatch.setattr(membrete, "DIR", tmp_path / "membretes")
    return tmp_path


@pytest.fixture
def _convertir_sin_libreoffice(monkeypatch):
    """`docx_a_pdf` falso: deja la hoja sintetica donde LibreOffice dejaria el PDF."""
    async def _falso(ruta_docx, carpeta):
        destino = Path(carpeta) / (Path(ruta_docx).stem + ".pdf")
        _hoja_pdf(destino)
        return str(destino), ""
    import shared.pdf_firma as pf
    monkeypatch.setattr(pf, "docx_a_pdf", _falso)


# ─── Medir ───────────────────────────────────────────────

def test_la_banda_termina_en_el_primer_hueco_y_tiene_que_empezar_cerca_del_borde():
    t, f = True, False
    assert membrete.banda([f, t, t, f, t, f, f, f, t], alcance=5, hueco=3) == 5
    # Lo que empieza despues del alcance no es membrete.
    assert membrete.banda([f, f, f, f, f, f, t, t], alcance=5, hueco=3) == 0
    assert membrete.banda([f] * 20, alcance=5, hueco=3) == 0


def test_se_mide_la_cabecera_y_el_pie_y_no_la_marca_de_agua(tmp_path):
    _hoja_pdf(tmp_path / "h.pdf", cabecera=110, pie=50)
    bandas = membrete.medir(tmp_path / "h.pdf")
    assert 108 <= bandas["cabecera_pt"] <= 112     # el bloque llega a 110 pt del borde
    assert 48 <= bandas["pie_pt"] <= 52


def test_una_hoja_sin_pie_mide_cero_abajo(tmp_path):
    _hoja_pdf(tmp_path / "h.pdf", pie=0)
    assert membrete.medir(tmp_path / "h.pdf")["pie_pt"] == 0


def test_los_margenes_dejan_aire_y_nunca_bajan_de_los_de_siempre():
    m = {"cabecera_pt": 143.0, "pie_pt": 0.0}
    arriba, abajo = membrete.margenes_pt(m, 56.7, 56.7)
    assert arriba == 143.0 + membrete.AIRE_PT
    assert abajo == 56.7
    assert membrete.margenes_pt(None, 56.7, 56.7) == (56.7, 56.7)


# ─── Guardar ─────────────────────────────────────────────

@pytest.mark.parametrize("contenido, motivo", [
    (b"", "vacío"),
    (b"%PDF-1.4 no soy un word", "Word"),
    (b"PK\x03\x04" + b"0" * 50, "Word"),
])
async def test_solo_se_acepta_un_word(_dir, contenido, motivo):
    with pytest.raises(membrete.MembreteInvalido, match=motivo):
        await membrete.guardar(5, contenido)
    assert membrete.de(5) is None


async def test_se_guarda_convertida_y_medida(_dir, _convertir_sin_libreoffice):
    bandas = await membrete.guardar(5, _hoja_docx())
    assert 108 <= bandas["cabecera_pt"] <= 112
    datos = membrete.de(5)
    assert datos["docx"].is_file() and datos["pdf"].is_file()
    assert json.loads((membrete.DIR / "empresa_5.json").read_text()) == bandas
    # Sin restos de la subida.
    assert sorted(p.name for p in membrete.DIR.iterdir()) == [
        "empresa_5.docx", "empresa_5.json", "empresa_5.pdf"]
    assert membrete.miniatura_png(5)[:8] == b"\x89PNG\r\n\x1a\n"

    membrete.borrar(5)
    assert membrete.de(5) is None and membrete.miniatura_png(5) is None


async def test_si_libreoffice_falla_no_queda_nada_a_medias(_dir, monkeypatch):
    async def _falla(ruta_docx, carpeta):
        return None, "LibreOffice no está instalado."
    import shared.pdf_firma as pf
    monkeypatch.setattr(pf, "docx_a_pdf", _falla)
    with pytest.raises(membrete.MembreteInvalido, match="LibreOffice"):
        await membrete.guardar(5, _hoja_docx())
    assert membrete.de(5) is None
    assert not any(membrete.DIR.iterdir())


# ─── Aplicarla ───────────────────────────────────────────

async def test_el_pdf_propio_sale_sobre_la_hoja_y_por_debajo_de_su_cabecera(
        _dir, _convertir_sin_libreoffice, monkeypatch):
    import shared.pdf_firmable as pf

    async def _sin_imagenes(_):
        return {}
    monkeypatch.setattr(pf, "rutas_de", _sin_imagenes)
    monkeypatch.setattr(pf, "SALIDA", _dir / "salida")
    await membrete.guardar(7, _hoja_docx())

    emp = {"id": 7, "razon_social": "EMPRESA DE PRUEBA S.A.C.", "ruc": "20000000001",
           "direccion": "Av. Prueba 1", "representante_legal": "TITULAR",
           "dni_representante": "00000000"}
    ruta = await pf.generar_pdf("con-membrete.pdf", "COTIZACIÓN", "REF-1",
                                ["Cuerpo de la cotización."] * 3, emp, con_dnie=True)

    hoja = PdfReader(ruta).pages[0]
    posiciones = {}

    def _anotar(texto, cm, tm, *_):
        if texto.strip():
            posiciones.setdefault(texto.strip(), tm[5])
    hoja.extract_text(visitor_text=_anotar)

    assert "MEMBRETE DE PRUEBA SAC  RUC 20000000001" in posiciones   # la hoja, de fondo
    techo = ALTO - membrete.de(7)["cabecera_pt"]
    assert posiciones["COTIZACIÓN"] < techo                          # debajo de la cabecera
    # Nuestra cabecera de logo + datos ya no se dibuja: la hoja lo dice.
    assert "RUC 20000000001" not in posiciones


async def test_sin_hoja_el_pdf_sigue_como_siempre(_dir, monkeypatch):
    import shared.pdf_firmable as pf

    async def _sin_imagenes(_):
        return {}
    monkeypatch.setattr(pf, "rutas_de", _sin_imagenes)
    monkeypatch.setattr(pf, "SALIDA", _dir / "salida")
    emp = {"id": 8, "razon_social": "OTRA S.A.C.", "ruc": "20000000002"}
    ruta = await pf.generar_pdf("sin.pdf", "COTIZACIÓN", "", ["x"], emp)
    assert "RUC 20000000002" in PdfReader(ruta).pages[0].extract_text()


async def test_el_word_sale_con_cabecera_cuerpo_vacio_y_margenes_sensatos(
        _dir, _convertir_sin_libreoffice):
    await membrete.guardar(9, _hoja_docx())
    doc = membrete.documento_word(9)
    assert doc.sections[0].header.paragraphs[0].text == "MEMBRETE DE PRUEBA SAC"
    assert [p.text for p in doc.paragraphs] == []
    assert doc.sections[0].left_margin >= Cm(2)
    assert doc.sections[0].top_margin >= Cm((membrete.de(9)["cabecera_pt"]
                                             + membrete.AIRE_PT) / 72 * 2.54) - 1

    # Y se puede escribir y guardar como cualquier Document().
    doc.add_paragraph("CARTA DE PRESENTACIÓN")
    salida = io.BytesIO()
    doc.save(salida)
    with zipfile.ZipFile(salida) as z:
        assert "CARTA DE PRESENTACIÓN" in z.read("word/document.xml").decode()


def test_sin_hoja_el_word_es_uno_en_blanco(_dir):
    doc = membrete.documento_word(10)
    assert doc.sections[0].header.paragraphs[0].text == ""


# ─── La ficha de la empresa ──────────────────────────────

async def _entrar(cliente, usuario):
    from shared.db import connection
    async with connection() as c:
        await c.execute("DELETE FROM intentos_acceso")
    r = await cliente.post("/entrar", data={"email": usuario["email"],
                                            "password": usuario["password"]})
    assert r.status_code == 303


@sin_base
async def test_subir_ver_y_quitar_la_hoja_desde_la_ficha(
        cliente, usuario, empresa, _dir, _convertir_sin_libreoffice):
    await _entrar(cliente, usuario)

    r = await cliente.post(f"/empresas/{empresa}/membrete",
                           files={"archivo": ("hoja.pdf", b"%PDF-1.4", "application/pdf")})
    assert "error=" in r.headers["location"]

    r = await cliente.post(f"/empresas/{empresa}/membrete",
                           files={"archivo": ("hoja.docx", _hoja_docx(),
                                              "application/octet-stream")})
    assert "aviso=" in r.headers["location"]

    ficha = await cliente.get(f"/empresas/{empresa}/editar")
    assert f"/empresas/{empresa}/membrete.png" in ficha.text
    vista = await cliente.get(f"/empresas/{empresa}/membrete.png")
    assert vista.status_code == 200 and vista.headers["content-type"] == "image/png"

    await cliente.post(f"/empresas/{empresa}/membrete/borrar")
    assert (await cliente.get(f"/empresas/{empresa}/membrete.png")).status_code == 404


@sin_base
async def test_la_hoja_de_otra_empresa_no_se_ve_ni_se_cambia(
        cliente, usuario, _dir, _convertir_sin_libreoffice):
    await membrete.guardar(999999, _hoja_docx())
    await _entrar(cliente, usuario)
    assert (await cliente.get("/empresas/999999/membrete.png")).status_code == 403
    r = await cliente.post("/empresas/999999/membrete/borrar")
    assert "no+es+tuya" in r.headers["location"]
    assert membrete.de(999999) is not None
